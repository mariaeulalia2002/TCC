import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def main(docx_path, out_dir):
    doc = Document(docx_path)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    records = []
    image_no = 0
    for p_index, paragraph in enumerate(doc.paragraphs):
        embeds = paragraph._p.xpath('.//a:blip/@r:embed')
        if not embeds:
            continue
        for rid in embeds:
            part = doc.part.related_parts[rid]
            image_no += 1
            suffix = Path(part.partname).suffix or '.bin'
            target = out / f'image-{image_no:02d}{suffix}'
            target.write_bytes(part.blob)
            before = doc.paragraphs[p_index - 1].text.strip() if p_index else ''
            after = doc.paragraphs[p_index + 1].text.strip() if p_index + 1 < len(doc.paragraphs) else ''
            records.append({
                'image_no': image_no,
                'paragraph_index': p_index,
                'file': str(target),
                'partname': str(part.partname),
                'before': before,
                'paragraph_text': paragraph.text.strip(),
                'after': after,
            })
    print(json.dumps(records, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
